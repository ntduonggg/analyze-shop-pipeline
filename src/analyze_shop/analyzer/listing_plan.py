"""Tab con "Kế hoạch listing" trong Build vs Conflict.

Target số listing MỚI mỗi shop build phải up tới PLAN_END, tách theo tuần và
dòng sản phẩm, kèm check conflict tương đối (% listing của shop conflict cùng
loại SP; >= 10% chặn, < 10% vẫn up được). Plan gốc sinh từ SHOP_PLAN bên dưới;
mọi chỉnh sửa trên trang được lưu vào PLAN_FILE trong thư mục build/conflict.
"""
import html
import io
import json
import math
import os
import re
from datetime import date, datetime, timedelta

import pandas as pd

PLAN_FILE = "listing_plan.json"
PLAN_START = date(2026, 10, 5)   # thứ 2 đầu tiên của plan
PLAN_END = date(2026, 12, 1)     # hạn chót phải có đủ listing
WORKDAYS_PER_WEEK = 5            # T2–T6; T7 chỉ làm 1 ngày/tháng (xem is_workday)
PER_DAY_MIN, PER_DAY_MAX = 15, 20  # năng lực tạo mẫu/ngày của cả team (1 mẫu = 1 listing)
BATCH = 5                        # mỗi lượt up = 5 listing cho 1 shop
# Số lượt (5 listing) mỗi ngày: ~10% số ngày làm việc chỉ 3 lượt (15 listing), dồn vào đầu W1 để
# mở dòng mới; còn lại 4 lượt (20 listing).
LOW_DAY_SHARE = 0.10
SHOPS_PER_DAY_LOW, SHOPS_PER_DAY_HIGH = 3, 4
# Số dòng SP mỗi shop tập trung (min, max); shop không có trong đây dùng LINES_DEFAULT.
LINES_DEFAULT = (4, 6)
LINES_PER_SHOP = {"YourStoryGiftsArt": (3, 3)}
DOW = ["T2", "T3", "T4", "T5", "T6", "T7", "CN"]
CONFLICT_BLOCK_PCT = 10
# Hệ số theo tuần: dòng Giáng Sinh dồn W1–W5 (listing mới cần 2–4 tuần mới có view),
# dòng evergreen tăng dần từ W6 để chuẩn bị Valentine/Q1.
SEASON_WEIGHT = {"X": [1.4, 1.4, 1.4, 1.4, 1.4, 1.0, 0.5, 0.5, 0.5],
                 "E": [0.8, 0.8, 0.8, 0.8, 0.8, 1.0, 1.4, 1.4, 1.4]}

# Apparel tách dòng con: shop conflict của shop clothing hầu hết bán 100% quần áo,
# so ở cấp "Apparel" thì dòng nào cũng bị chặn.
APPAREL_RULES = [
    ("Baby Bodysuit/Onesie", r"onesie|bodysuit|romper|baby (?:shirt|tee)|newborn outfit"),
    ("Pajamas", r"pajama|pjs|sleepwear|loungewear|nightwear"),
    ("Hoodie", r"hoodie"),
    ("Sweatshirt", r"sweatshirt|crewneck|sweater"),
    ("Kids/Toddler Tee", r"(?=.*(?:toddler|kids?|youth|child))(?=.*(?:shirt|tee))"),
    ("T-Shirt", r"shirt|\btee\b"),
    ("Hat/Beanie", r"\bhat\b|beanie|cap\b"),
    ("Socks", r"socks"),
    ("Tank top", r"tank top"),
    ("Dress/Skirt", r"dress|skirt|skort"),
    ("Jacket", r"jacket"),
    ("Costume", r"costume"),
    ("Swim", r"swim"),
    ("Leggings/Activewear", r"legging|activewear"),
]
_APPAREL_RES = [(n, re.compile(p)) for n, p in APPAREL_RULES]

# Loại SP cân nhắc theo từng hướng — dùng để liệt kê "không up (conflict >= 10%)".
DIRECTION_CANDIDATES = {
    "Handmade": ["Embroidered Ornament", "Stocking", "Embroidered Wreath Sash", "Embroidered Baby Photo Album",
                 "Embroidered Name Hoop", "Embroidered Name Pennant", "Embroidered Pennant Bag Tag",
                 "Embroidered Wedding Guest Book", "Embroidered Wedding Pennant Flag", "Stained Glass Suncatcher",
                 "Book Vase", "Wooden Key Holder", "Linen/Fabric Sign", "Acrylic Ornament", "Ceramic Ornament",
                 "Personalized Jewelry Dish", "Pillow"],
    "Clothing": [n for n, _ in APPAREL_RULES],
    "Custom Gift POD": ["Acrylic Ornament", "Ceramic Ornament", "Glass Ornament", "Wooden Ornament 1 Layer",
                        "Ceramic Mug", "Frosted Cup", "Tumbler", "Acrylic Magnet", "Cruise Door Magnet",
                        "Acrylic Keychain", "Custom Badge Reel", "Acrylic Plaque/Block", "Pillow", "Blanket",
                        "Wall Tapestry", "Beach Towel", "Doormat", "Custom Can Cooler", "Night Light/Lamp",
                        "Sticker/Decal", "Car Accessory", "Personalized Compact Mirror", "Personalized Jewelry Dish",
                        "Wooden Sign 1 Layer", "Stocking", "Reusable Advent Calendar", "Linen/Fabric Sign",
                        "Wedding Napkins"],
}

# Plan gốc: shop -> (hướng, % công suất, [(dòng SP, loại SP, % trong shop, mùa vụ, ghi chú, tuần bắt đầu)]).
# Mùa vụ: "X" = Giáng Sinh, "E" = evergreen. Hướng "Clothing" so conflict theo APPAREL_RULES.
# Mỗi shop tập trung 4–6 dòng (LINES_PER_SHOP), YourStoryGiftsArt chỉ 3 dòng thế mạnh.
SHOP_PLAN = {
    "VowgiftNina": ("Handmade", 12, [
        ("Embroidered Ornament", "Embroidered Ornament", 26, "X", "Thế mạnh hiện tại. Ornament thêu tên/năm/pet", 1),
        ("Embroidered Christmas Stocking", "Stocking", 12, "X", "Stocking thêu tên, ra sớm trước 15/11", 1),
        ("Embroidered Baby Photo Album", "Embroidered Baby Photo Album", 19, "E", "Nhu cầu cao, ~99 views/tháng/listing", 1),
        ("Embroidered Name Hoop", "Embroidered Name Hoop", 17, "E", "Hoop tên bé / cặp đôi / pet", 1),
        ("Embroidered Name Pennant", "Embroidered Name Pennant", 13, "E", "", 1),
        ("Embroidered Pennant Bag Tag", "Embroidered Pennant Bag Tag", 13, "E", "Giá thấp, dễ bán kèm", 1),
    ]),
    "YourStoryGiftsArt": ("Clothing", 18, [
        ("Sweatshirt (ugly/funny Christmas, custom photo)", "Sweatshirt", 55, "X", "Thế mạnh hiện tại của shop", 1),
        ("Hoodie", "Hoodie", 28, "X", "Tránh mẫu Disney/baby tee của UnboxJoyStudio", 1),
        ("Custom Face/Pet Socks", "Socks", 17, "X", "Gần như chưa ai làm trong data", 1),
    ]),
    "WonderSeedCo": ("Custom Gift POD", 14, [
        ("Night Light/Lamp", "Night Light/Lamp", 27, "X", "Thế mạnh: ~36 views/tháng/listing", 1),
        ("Wooden Ornament 1 Layer (UV print)", "Wooden Ornament 1 Layer", 24, "X", "Ornament còn mở cho shop", 1),
        ("Blanket", "Blanket", 17, "X", "", 1),
        ("Tumbler", "Tumbler", 16, "E", "Nhu cầu cao", 1),
        ("Custom Badge Reel", "Custom Badge Reel", 16, "E", "Nhu cầu cao nhất data", 1),
    ]),
    "MapleGiftCo": ("Custom Gift POD", 14, [
        ("Glass Ornament", "Glass Ornament", 22, "X", "Ornament còn mở (acrylic/ceramic bị chặn)", 1),
        ("Pillow", "Pillow", 16, "X", "Đang bán", 1),
        ("Blanket", "Blanket", 14, "X", "", 1),
        ("Ceramic Mug", "Ceramic Mug", 17, "E", "Đang bán", 1),
        ("Acrylic Magnet", "Acrylic Magnet", 17, "E", "Đang bán", 1),
        ("Frosted Cup", "Frosted Cup", 14, "E", "Đang bán", 1),
    ]),
    "MontiKasGift": ("Custom Gift POD", 14, [
        ("Ceramic Ornament", "Ceramic Ornament", 27, "X", "Đang bán", 1),
        ("Acrylic Ornament", "Acrylic Ornament", 17, "X", "Chờ VowgiftNina (đang 13.2%) xuống dưới 10% khi up handmade", 4),
        ("Personalized Jewelry Dish", "Personalized Jewelry Dish", 26, "E", "Thế mạnh hiện tại", 1),
        ("Tumbler", "Tumbler", 15, "E", "", 1),
        ("Car Accessory", "Car Accessory", 15, "E", "Nhu cầu cao", 1),
    ]),
    "NatureStoreAU": ("Custom Gift POD", 14, [
        ("Acrylic Ornament", "Acrylic Ornament", 29, "X", "", 1),
        ("Glass Ornament", "Glass Ornament", 16, "X", "", 1),
        ("Custom Badge Reel", "Custom Badge Reel", 23, "E", "Nhu cầu cao nhất data, shop đã có 2 listing", 1),
        ("Ceramic Mug", "Ceramic Mug", 16, "E", "", 1),
        ("Tumbler", "Tumbler", 16, "E", "", 1),
    ]),
    "JoyMadeCoGift": ("Custom Gift POD", 14, [
        ("Ceramic Ornament", "Ceramic Ornament", 26, "X", "Đang bán", 1),
        ("Acrylic Ornament", "Acrylic Ornament", 23, "X", "Đang bán", 1),
        ("Glass Ornament", "Glass Ornament", 13, "X", "Đang bán", 1),
        ("Tumbler", "Tumbler", 21, "E", "", 1),
        ("Ceramic Mug", "Ceramic Mug", 17, "E", "", 1),
    ]),
}


def is_workday(d: date) -> bool:
    """T2–T6, cộng thêm T7 đầu tiên của tháng (T7 của tuần đầu tháng)."""
    return d.weekday() < WORKDAYS_PER_WEEK or (d.weekday() == 5 and d.day <= 7)


def plan_weeks() -> list:
    """Tuần T2–CN từ PLAN_START tới PLAN_END (tuần cuối có thể ngắn), chỉ tính ngày làm
    việc theo is_workday. "end" là CN để mọi ngày đều thuộc 1 tuần; "to" là ngày làm cuối."""
    n_all = sum(is_workday(PLAN_START + timedelta(days=i)) for i in range((PLAN_END - PLAN_START).days + 1))
    n_low = round(n_all * LOW_DAY_SHARE)
    out, mon, k, seen = [], PLAN_START, 1, 0
    while mon <= PLAN_END:
        end = min(mon + timedelta(days=6), PLAN_END)
        dates = [mon + timedelta(days=i) for i in range((end - mon).days + 1) if is_workday(mon + timedelta(days=i))]
        caps = [SHOPS_PER_DAY_LOW if seen + j < n_low else SHOPS_PER_DAY_HIGH for j in range(len(dates))]
        seen += len(dates)
        slots = sum(caps)
        out.append({"code": f"W{k}", "from": mon.strftime("%d/%m"), "to": dates[-1].strftime("%d/%m"),
                    "start": mon.isoformat(), "end": end.isoformat(), "days": len(dates),
                    "dates": [{"iso": d.isoformat(), "label": f"{DOW[d.weekday()]} {d:%d/%m}", "cap": c} for d, c in zip(dates, caps)],
                    "slots": slots, "total": slots * BATCH, "per_day": round(slots * BATCH / len(dates), 1)})
        mon += timedelta(days=7)
        k += 1
    return out


def apparel_sub(title: str) -> str:
    t = html.unescape(str(title or "")).lower()
    for name, rx in _APPAREL_RES:
        if rx.search(t):
            return name
    return "Apparel khác"


def _shop_slots(weeks: list, shops: list) -> dict:
    """Số lượt up (5 listing/lượt) mỗi shop mỗi tuần theo % công suất. Chia cộng dồn
    để phần lẻ xoay vòng giữa các tuần; mỗi shop tối đa 1 lượt/ngày."""
    wt = {s: SHOP_PLAN[s][1] for s in shops}
    tot_w = sum(wt.values())
    done = {s: 0 for s in shops}
    out = {s: [0] * len(weeks) for s in shops}
    cum = 0
    for w, wk in enumerate(weeks):
        cum += wk["slots"]
        alloc = {s: 0 for s in shops}
        for _ in range(wk["slots"]):
            ok = [s for s in shops if alloc[s] < wk["days"]]
            s = max(ok, key=lambda s: (wt[s] / tot_w * cum - done[s] - alloc[s], wt[s]))
            alloc[s] += 1
        for s in shops:
            out[s][w] = alloc[s]
            done[s] += alloc[s]
    return out


def _line_batches(rows: list) -> list:
    """Lượt 5 listing của 1 shop trong tuần từ target từng dòng: [(row, số listing)].
    Target không chia hết cho 5 thì lượt cuối của dòng đó ít hơn 5. Xếp xen kẽ các dòng
    (round-robin) để 1 dòng nhiều lượt không dồn hết vào đầu tuần."""
    q = [[(r, min(BATCH, r["t"] - k)) for k in range(0, r["t"], BATCH)] for r in rows if r["t"] > 0]
    out = []
    while any(q):
        for b in q:
            if b:
                out.append(b.pop(0))
    return out


def daily_schedule(week: dict, week_rows: dict, order: list) -> list:
    """Chia target tuần thành lượt up theo ngày. Mỗi lượt = 5 listing của 1 dòng SP cho
    1 shop; mỗi shop tối đa 1 lượt/ngày, số lượt mỗi ngày theo công suất ngày (3 hoặc 4). Nhiều shop
    cùng ngày được up cùng dòng SP (dùng chung mẫu) — khi chọn lượt cho shop sẽ ưu tiên
    dòng SP đã có shop khác up trong ngày. week_rows: shop -> [{"line","ptype","t"}].
    Trả về [{"label", "iso", "items": [(shop, dòng SP, loại SP, số listing)]}]."""
    queue = {s: _line_batches(week_rows.get(s, [])) for s in order}
    need = {s: len(queue[s]) for s in order}
    days = week["dates"]
    out = []
    for i, d in enumerate(days):
        left = len(days) - i
        # Số lượt hôm nay = công suất ngày (cap); chỉ vượt khi các ngày sau không gánh hết phần còn lại.
        later = sum(x.get("cap", 0) for x in days[i + 1:])
        cap = max(d.get("cap", 0), sum(need.values()) - later, math.ceil(sum(need.values()) / left) if not d.get("cap") else 0)
        must = [s for s in order if need[s] >= left]  # còn nhiều lượt hơn số ngày -> phải up hôm nay
        pick = sorted([s for s in order if need[s] > 0], key=lambda s: (-need[s], order.index(s)))
        pick = must + [s for s in pick if s not in must][:max(0, cap - len(must))]
        items, today = [], set()
        for s in sorted(pick, key=order.index):
            q = queue[s]
            j = next((k for k, (r, _) in enumerate(q) if r["ptype"] in today), 0)
            r, n = q.pop(j)
            need[s] -= 1
            today.add(r["ptype"])
            items.append((s, r["line"], r["ptype"], n))
        out.append({"label": d["label"], "iso": d["iso"], "items": items})
    return out


def default_rows(weeks: list) -> list:
    """Plan gốc: lượt mỗi shop mỗi tuần (_shop_slots) chia cho các dòng SP theo % trong
    shop × hệ số mùa vụ. Chia cộng dồn qua các tuần — mỗi lượt giao cho dòng đang thiếu
    nhiều nhất so với tỷ lệ — nên dòng nhỏ vẫn tới lượt, target mỗi ô luôn là bội số của 5."""
    nw = len(weeks)
    shops = list(SHOP_PLAN)
    slots = _shop_slots(weeks, shops)
    rows = []
    for s in shops:
        direction, _, lines = SHOP_PLAN[s]
        grid = [[0] * nw for _ in lines]
        want, got = [0.0] * len(lines), [0] * len(lines)
        want_g, got_g = {"X": 0.0, "E": 0.0}, {"X": 0, "E": 0}
        for w in range(nw):
            ws = [l[2] * SEASON_WEIGHT[l[3]][min(w, len(SEASON_WEIGHT[l[3]]) - 1)] * (w + 1 >= l[5]) for l in lines]
            tot = sum(ws)
            if not tot:
                continue
            for i, x in enumerate(ws):
                want[i] += slots[s][w] * x / tot
                want_g[lines[i][3]] += slots[s][w] * x / tot
            # Chọn mùa vụ (Giáng Sinh/evergreen) trước để tỷ lệ mỗi tuần đều, rồi chọn dòng trong mùa đó.
            for _ in range(slots[s][w]):
                g = max((g for g in want_g if any(ws[i] > 0 and lines[i][3] == g for i in range(len(lines)))),
                        key=lambda g: want_g[g] - got_g[g])
                i = max((i for i in range(len(lines)) if ws[i] > 0 and lines[i][3] == g), key=lambda i: (want[i] - got[i], ws[i]))
                got[i] += 1
                got_g[g] += 1
                grid[i][w] += BATCH
        for i, (name, ptype, _, season, note, start) in enumerate(lines):
            rows.append({"id": f"{s}-{i + 1:02d}", "shop": s, "order": i + 1, "line": name, "ptype": ptype,
                         "apparel": direction == "Clothing", "season": season, "start": start,
                         "note": note, "t": grid[i], "d": [0] * nw})
    return rows


def load_rows(data_dir: str, weeks: list):
    """(rows, saved_at) — plan đã lưu, hoặc plan gốc nếu chưa lưu / file hỏng."""
    path = os.path.join(data_dir, PLAN_FILE)
    nw = len(weeks)
    if os.path.isfile(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                saved = json.load(f)
            rows = saved["rows"]
            for r in rows:  # số tuần đổi (sửa PLAN_END) -> cắt/thêm cột cho khớp
                r["t"] = (list(r.get("t", [])) + [0] * nw)[:nw]
                r["d"] = (list(r.get("d", [])) + [0] * nw)[:nw]
            return rows, saved.get("saved_at")
        except Exception as e:
            print(f"[warn] Không đọc được {path}: {e} — dùng plan gốc")
    return default_rows(weeks), None


def _clean_int(v) -> int:
    try:
        return max(0, min(999, int(v)))
    except (TypeError, ValueError):
        return 0


def save_rows(data_dir: str, rows: list, nw: int) -> str:
    clean = []
    for r in rows:
        clean.append({
            "id": str(r["id"])[:80], "shop": str(r["shop"])[:80], "order": _clean_int(r.get("order")),
            "line": str(r.get("line", ""))[:200], "ptype": str(r.get("ptype", ""))[:80],
            "apparel": bool(r.get("apparel")), "season": "X" if r.get("season") == "X" else "E",
            "start": max(1, min(nw, _clean_int(r.get("start")) or 1)), "note": str(r.get("note", ""))[:500],
            "t": [_clean_int(x) for x in (list(r.get("t", [])) + [0] * nw)[:nw]],
            "d": [_clean_int(x) for x in (list(r.get("d", [])) + [0] * nw)[:nw]],
        })
    saved_at = datetime.now().isoformat(timespec="seconds")
    path = os.path.join(data_dir, PLAN_FILE)
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump({"saved_at": saved_at, "rows": clean}, f, ensure_ascii=False, indent=1)
    os.replace(tmp, path)
    return saved_at


def reset_rows(data_dir: str):
    path = os.path.join(data_dir, PLAN_FILE)
    if os.path.isfile(path):
        os.replace(path, path + f".bak-{datetime.now():%Y%m%d_%H%M%S}")


def build_context(listings: pd.DataFrame, build_list: list) -> dict:
    """Số liệu để tính conflict: số listing mỗi loại SP (và dòng con Apparel) của
    từng shop, tổng listing mỗi shop. listings đã có cột _ptype."""
    df = listings.copy()
    df["_shop"] = df["shop_name"].astype(str).str.lower()
    ptype_n = {s: {k: int(v) for k, v in g["_ptype"].value_counts().items()} for s, g in df.groupby("_shop")}
    app = df[df["_ptype"] == "Apparel"].copy()
    app["_sub"] = app["title"].map(apparel_sub)
    app_n = {s: {k: int(v) for k, v in g["_sub"].value_counts().items()} for s, g in app.groupby("_shop")}
    total = {s: int(n) for s, n in df.groupby("_shop").size().items()}
    return {"ptype_n": ptype_n, "app_n": app_n, "total": total,
            "conflicts": {b["build_shop"]: b["conflicts"] for b in build_list}}


def row_conflict(row: dict, rows: list, ctx: dict) -> dict:
    """Tỷ lệ % listing cùng loại SP của từng shop conflict TẠI TUẦN BẮT ĐẦU của dòng.
    Shop conflict nào cũng là shop build thì cộng thêm listing nó sẽ up trước tuần đó
    theo plan (vd VowgiftNina up handmade làm % Acrylic Ornament giảm dần).
    block: >= 10% tại tuần bắt đầu · wait: hiện >= 10% nhưng xuống dưới 10% tại tuần
    bắt đầu · light: > 0% · ok: không shop conflict nào bán."""
    key, app, start = row["ptype"], bool(row.get("apparel")), int(row.get("start") or 1)
    src = ctx["app_n"] if app else ctx["ptype_n"]
    hits, max_now = [], 0.0
    for c in ctx["conflicts"].get(row["shop"], []):
        ck = c.lower()
        base_n, base_tot = src.get(ck, {}).get(key, 0), ctx["total"].get(ck, 0)
        c_rows = [r for r in rows if r["shop"].lower() == ck]
        add_tot = sum(sum(r["t"][:start - 1]) for r in c_rows)
        add_n = sum(sum(r["t"][:start - 1]) for r in c_rows if r["ptype"] == key and bool(r.get("apparel")) == app)
        planned = any(sum(r["t"]) > 0 for r in c_rows if r["ptype"] == key and bool(r.get("apparel")) == app)
        now = round(base_n / base_tot * 100, 1) if base_tot else 0.0
        at = round((base_n + add_n) / (base_tot + add_tot) * 100, 1) if base_tot + add_tot else 0.0
        max_now = max(max_now, now)
        if at > 0 or now > 0 or planned:
            hits.append({"shop": c, "pct": at, "now": now, "planned": planned})
    hits.sort(key=lambda h: -h["pct"])
    mx = hits[0]["pct"] if hits else 0.0
    if mx >= CONFLICT_BLOCK_PCT:
        level = "block"
    elif max_now >= CONFLICT_BLOCK_PCT:
        level = "wait"
    elif mx > 0 or any(h["planned"] for h in hits):
        level = "light"
    else:
        level = "ok"
    return {"level": level, "max": mx, "hits": hits}


def build_payload(listings: pd.DataFrame, build_list: list, profile, market: dict, data_dir: str) -> dict:
    weeks = plan_weeks()
    ctx = build_context(listings, build_list)
    rows, saved_at = load_rows(data_dir, weeks)
    shops = []
    for b in build_list:
        s = b["build_shop"]
        p = profile(s) or {}
        direction = SHOP_PLAN[s][0] if s in SHOP_PLAN else "Custom Gift POD"
        app = direction == "Clothing"
        blocked = []
        for k in DIRECTION_CANDIDATES[direction]:
            c = row_conflict({"shop": s, "ptype": k, "apparel": app, "start": 1}, [], ctx)
            if c["level"] == "block":
                blocked.append({"line": k, "max": c["max"], "hits": c["hits"]})
        shops.append({"shop": s, "direction": direction, "current": p.get("n_listings", 0),
                      "lines": list(LINES_PER_SHOP.get(s, LINES_DEFAULT)),
                      "conflicts": b["conflicts"],
                      "current_mix": [{"ptype": d["ptype"], "pct": d["pct"]} for d in p.get("ptype_detail", [])[:3]],
                      "blocked": blocked})
    return {
        "weeks": weeks, "shops": shops, "rows": rows, "saved_at": saved_at,
        "ctx": {k: ctx[k] for k in ("ptype_n", "app_n", "total")},
        "apparel_rules": APPAREL_RULES, "block_pct": CONFLICT_BLOCK_PCT,
        "per_day_min": PER_DAY_MIN, "per_day_max": PER_DAY_MAX, "batch": BATCH,
        "demand": {k: v.get("demand") for k, v in market.items() if v.get("demand") is not None},
        "plan_file": os.path.join(data_dir, PLAN_FILE),
    }


# ── Xuất Excel ───────────────────────────────────────────────────────────────
LEVEL_TXT = {"ok": "OK", "light": "Nhẹ (<10%)", "wait": "Chờ mở", "block": "Chặn (≥10%)"}
LEVEL_FILL = {"OK": "E2F0D9", "Nhẹ (<10%)": "FFF2CC", "Chờ mở": "DDEBF7", "Chặn (≥10%)": "F8CBAD"}


def plan_xlsx(payload: dict, ctx_conflicts: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as CL

    W, rows, shops = payload["weeks"], payload["rows"], payload["shops"]
    NW = len(W)
    ctx = dict(payload["ctx"], conflicts=ctx_conflicts)
    shop_order = [s["shop"] for s in shops]
    rows = sorted(rows, key=lambda r: (shop_order.index(r["shop"]) if r["shop"] in shop_order else 99, r["order"]))
    dirs = {s["shop"]: s["direction"] for s in shops}
    F = "Arial"
    H = Font(name=F, bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="2F4858")
    B, BB = Font(name=F), Font(name=F, bold=True)
    BLUE, GREEN = Font(name=F, color="0000FF"), Font(name=F, color="008000")
    ITAL = Font(name=F, italic=True, size=9)
    YEL, GREY = PatternFill("solid", fgColor="FFF2CC"), PatternFill("solid", fgColor="F2F2F2")
    thin = Side(style="thin", color="D0D0D0"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
    CEN = Alignment(horizontal="center")

    def header(ws, cols, widths):
        for i, c in enumerate(cols, 1):
            x = ws.cell(row=1, column=i, value=c); x.font = H; x.fill = HF; x.border = BOX
            x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 32
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[CL(i)].width = w

    def cell(ws, r, c, v, font=B, fill=None, fmt=None, align=None):
        x = ws.cell(row=r, column=c, value=v); x.font = font; x.border = BOX
        if fill: x.fill = fill
        if fmt: x.number_format = fmt
        if align: x.alignment = align
        return x

    def hit_txt(c):
        return ", ".join(f'{h["shop"]} {h["pct"]}%' + (" (đã lên plan)" if h["planned"] else "") for h in c["hits"]) or "—"

    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    wk_hdr = [f'{w["code"]}\n{w["from"]}–{w["to"]}' for w in W]
    last = len(rows) + 1; tr = last + 1

    # Kế hoạch tuần
    kp = wb.active; kp.title = "Kế hoạch tuần"
    WC0 = 10
    header(kp, ["ID", "Shop", "Hướng", "Dòng sản phẩm", "Loại SP", "Mùa vụ", "Conflict", "Max % conflict",
                "Shop conflict (% tại tuần bắt đầu)"] + wk_hdr + ["Tổng target", "Tuần bắt đầu", "Ghi chú"],
           [16, 18, 15, 38, 24, 11, 13, 10, 36] + [8] * NW + [10, 9, 48])
    for i, r in enumerate(rows, 2):
        c = row_conflict(r, rows, ctx)
        vals = [r["id"], r["shop"], dirs.get(r["shop"], ""), r["line"], r["ptype"],
                "Giáng Sinh" if r["season"] == "X" else "Evergreen", LEVEL_TXT[c["level"]], c["max"] / 100, hit_txt(c)]
        for j, v in enumerate(vals, 1):
            cell(kp, i, j, v, fmt="0.0%" if j == 8 else None)
        kp.cell(row=i, column=7).fill = PatternFill("solid", fgColor=LEVEL_FILL[LEVEL_TXT[c["level"]]])
        for w in range(NW):
            cell(kp, i, WC0 + w, r["t"][w], font=BLUE, fill=YEL, align=CEN)
        cell(kp, i, WC0 + NW, f"=SUM({CL(WC0)}{i}:{CL(WC0 + NW - 1)}{i})", font=BB, align=CEN)
        cell(kp, i, WC0 + NW + 1, f'W{r["start"]}', align=CEN)
        cell(kp, i, WC0 + NW + 2, r.get("note", ""))
    cell(kp, tr, 4, "TỔNG", font=BB, fill=GREY)
    for w in range(NW + 1):
        col = CL(WC0 + w)
        cell(kp, tr, WC0 + w, f"=SUM({col}2:{col}{last})", font=BB, fill=GREY, align=CEN)
    kp.freeze_panes = "E2"; kp.auto_filter.ref = f"A1:{CL(WC0 + NW + 2)}{last}"
    TOTC = CL(WC0 + NW)

    # Theo dõi thực tế
    td = wb.create_sheet("Theo dõi thực tế")
    header(td, ["ID", "Shop", "Dòng sản phẩm"] + wk_hdr + ["Đã up", "Target", "% hoàn thành", "Còn thiếu"],
           [16, 18, 38] + [8] * NW + [9, 9, 12, 10])
    DU, TG, PC, MS = CL(4 + NW), CL(5 + NW), CL(6 + NW), CL(7 + NW)
    for i, r in enumerate(rows, 2):
        for j, col in enumerate("ABD", 1):
            cell(td, i, j, f"='Kế hoạch tuần'!{col}{i}", font=GREEN)
        for w in range(NW):
            cell(td, i, 4 + w, r["d"][w] or None, font=BLUE, fill=YEL, align=CEN)
        cell(td, i, 4 + NW, f"=SUM({CL(4)}{i}:{CL(3 + NW)}{i})", font=BB, align=CEN)
        cell(td, i, 5 + NW, f"='Kế hoạch tuần'!{TOTC}{i}", font=GREEN, align=CEN)
        cell(td, i, 6 + NW, f"=IF({TG}{i}=0,0,{DU}{i}/{TG}{i})", fmt="0%", align=CEN)
        cell(td, i, 7 + NW, f"=MAX(0,{TG}{i}-{DU}{i})", align=CEN)
    cell(td, tr, 3, "TỔNG", font=BB, fill=GREY)
    for c in range(4, 6 + NW):
        cell(td, tr, c, f"=SUM({CL(c)}2:{CL(c)}{last})", font=BB, fill=GREY, align=CEN)
    cell(td, tr, 6 + NW, f"=IF({TG}{tr}=0,0,{DU}{tr}/{TG}{tr})", font=BB, fill=GREY, fmt="0%")
    cell(td, tr, 7 + NW, f"=MAX(0,{TG}{tr}-{DU}{tr})", font=BB, fill=GREY)
    td.freeze_panes = "D2"
    td.conditional_formatting.add(f"{PC}2:{PC}{tr}", CellIsRule(operator="greaterThanOrEqual", formula=["1"],
                                                                fill=PatternFill("solid", fgColor="E2F0D9")))
    td.cell(row=tr + 2, column=1, value="Ô vàng chữ xanh = số listing đã up thực tế mỗi tuần.").font = ITAL

    # Công suất tuần
    cs = wb.create_sheet("Công suất tuần")
    header(cs, ["Tuần", "Từ", "Đến", "Ngày làm việc", "Listing/ngày (TB kế hoạch)", "Target tuần", f"Min ({PER_DAY_MIN}/ngày)",
                f"Max ({PER_DAY_MAX}/ngày)", "Đã phân cho dòng SP", "Lệch (phân − target)", "Trong ngưỡng?"],
           [8, 8, 8, 10, 12, 11, 12, 12, 13, 13, 12])
    pmin, pmax = NW + 5, NW + 6
    for i, w in enumerate(W, 2):
        for j, v in enumerate([w["code"], w["from"], w["to"]], 1):
            cell(cs, i, j, v)
        cell(cs, i, 4, w["days"], font=BLUE, fill=YEL); cell(cs, i, 5, w["per_day"], font=BLUE, fill=YEL)
        cell(cs, i, 6, f"=D{i}*E{i}", font=BB)
        cell(cs, i, 7, f"=D{i}*$B${pmin}"); cell(cs, i, 8, f"=D{i}*$B${pmax}")
        cell(cs, i, 9, f"='Kế hoạch tuần'!{CL(WC0 + i - 2)}{tr}", font=GREEN)
        cell(cs, i, 10, f"=I{i}-F{i}")
        cell(cs, i, 11, f'=IF(AND(I{i}>=G{i},I{i}<=H{i}),"Có","Không")')
    t = NW + 2
    cell(cs, t, 1, "TỔNG", font=BB, fill=GREY)
    for c in [4, 6, 7, 8, 9, 10]:
        cell(cs, t, c, f"=SUM({CL(c)}2:{CL(c)}{t - 1})", font=BB, fill=GREY)
    cell(cs, pmin, 1, "Min mẫu/ngày", font=BB); cell(cs, pmin, 2, PER_DAY_MIN, font=BLUE, fill=YEL)
    cell(cs, pmax, 1, "Max mẫu/ngày", font=BB); cell(cs, pmax, 2, PER_DAY_MAX, font=BLUE, fill=YEL)
    cs.cell(row=pmin, column=3, value=f"Mỗi lượt = 5 listing của 1 dòng SP cho 1 shop. ~{LOW_DAY_SHARE:.0%} số ngày {SHOPS_PER_DAY_LOW} lượt ({SHOPS_PER_DAY_LOW * BATCH} listing, đầu W1), "
            f"còn lại {SHOPS_PER_DAY_HIGH} lượt ({SHOPS_PER_DAY_HIGH * BATCH} listing). Làm T2–T6, thêm T7 đầu tiên của tháng.").font = ITAL
    cs.conditional_formatting.add(f"K2:K{NW + 1}", CellIsRule(operator="equal", formula=['"Không"'],
                                                              fill=PatternFill("solid", fgColor="F8CBAD")))

    # Shop x Tuần
    sx = wb.create_sheet("Shop x Tuần")
    header(sx, ["Shop", "Hướng"] + [w["code"] for w in W] + ["Tổng target"], [20, 16] + [8] * NW + [11])
    for i, s in enumerate(shops, 2):
        cell(sx, i, 1, s["shop"], font=BB); cell(sx, i, 2, s["direction"])
        for w in range(NW):
            wc = CL(WC0 + w)
            cell(sx, i, 3 + w, f"=SUMIF('Kế hoạch tuần'!$B$2:$B${last},$A{i},'Kế hoạch tuần'!{wc}$2:{wc}${last})", align=CEN)
        cell(sx, i, 3 + NW, f"=SUM({CL(3)}{i}:{CL(2 + NW)}{i})", font=BB, align=CEN)
    st = len(shops) + 2
    cell(sx, st, 1, "TỔNG", font=BB, fill=GREY)
    for c in range(3, 4 + NW):
        cell(sx, st, c, f"=SUM({CL(c)}2:{CL(c)}{st - 1})", font=BB, fill=GREY, align=CEN)

    # Lịch up theo ngày: mỗi dòng = 1 lượt 5 listing của 1 dòng SP cho 1 shop
    ld = wb.create_sheet("Lịch up theo ngày")
    header(ld, ["Tuần", "Ngày", "Shop", "Dòng sản phẩm", "Loại SP", "Số listing", "Đã up", "Dùng chung mẫu với"],
           [7, 11, 20, 40, 24, 10, 9, 36])
    ri = 2
    B5 = PatternFill("solid", fgColor="E2EFDA")
    for wi, w in enumerate(W):
        wr = {s: [dict(r, t=r["t"][wi]) for r in rows if r["shop"] == s] for s in shop_order}
        a = ri
        for d in daily_schedule(w, wr, shop_order):
            d0 = ri
            for s, line, ptype, n in d["items"]:
                same = [x[0] for x in d["items"] if x[2] == ptype and x[0] != s]
                for j, v in enumerate([w["code"], d["label"], s, line, ptype], 1):
                    cell(ld, ri, j, v, font=BB if j == 2 else B)
                cell(ld, ri, 6, n, font=BB, fill=B5 if n == BATCH else YEL, align=CEN)
                cell(ld, ri, 7, None, font=BLUE, fill=YEL, align=CEN)
                cell(ld, ri, 8, ", ".join(same) or None)
                ri += 1
            cell(ld, ri, 2, d["label"], font=ITAL, fill=GREY)
            cell(ld, ri, 3, f'=COUNTA(C{d0}:C{ri - 1})&" shop"', font=ITAL, fill=GREY)
            for c in (4, 5, 8):
                cell(ld, ri, c, None, fill=GREY)
            cell(ld, ri, 6, f"=SUM(F{d0}:F{ri - 1})", font=BB, fill=GREY, align=CEN)
            cell(ld, ri, 7, f"=SUM(G{d0}:G{ri - 1})", font=BB, fill=GREY, align=CEN)
            ri += 1
        cell(ld, ri, 1, w["code"], font=BB, fill=GREY); cell(ld, ri, 2, f"Tổng {w['code']}", font=BB, fill=GREY)
        for c in (3, 4, 5, 8):
            cell(ld, ri, c, None, fill=GREY)
        cell(ld, ri, 6, f'=SUMIFS(F{a}:F{ri - 1},C{a}:C{ri - 1},"<>*shop")', font=BB, fill=GREY, align=CEN)
        cell(ld, ri, 7, f'=SUMIFS(G{a}:G{ri - 1},C{a}:C{ri - 1},"<>*shop")', font=BB, fill=GREY, align=CEN)
        ri += 2
    ld.freeze_panes = "D2"
    ld.cell(row=ri, column=1, value=f"Mỗi dòng = 1 lượt {BATCH} listing của 1 dòng SP cho 1 shop; mỗi shop tối đa 1 lượt/ngày. "
            "Nhiều shop cùng ngày có thể up cùng dòng SP (cột 'Dùng chung mẫu với'). Ô vàng ở cột Số listing = lượt "
            "thiếu vì target không chia hết cho 5. Lịch tự chia từ target tuần ở sheet 'Kế hoạch tuần' khi xuất file.").font = ITAL

    # Tổng quan (sheet đầu)
    tq = wb.create_sheet("Tổng quan", 0)
    tq["A1"] = f"Kế hoạch listing shop build — {W[0]['from']} → {PLAN_END:%d/%m/%Y}"; tq["A1"].font = Font(name=F, bold=True, size=14)
    tq["A2"] = f"Xuất từ tab Build vs Conflict › Kế hoạch listing lúc {datetime.now():%d/%m/%Y %H:%M}."; tq["A2"].font = ITAL
    hdr = ["Shop", "Hướng", "Listing hiện có", f"Target mới đến {PLAN_END:%d/%m}", f"Tổng listing {PLAN_END:%d/%m}",
           "Đã up (thực tế)", "% hoàn thành", "Số dòng SP", "Shop conflict", "Đang bán nhiều nhất"]
    for i, c in enumerate(hdr, 1):
        x = tq.cell(row=4, column=i, value=c); x.font = H; x.fill = HF; x.border = BOX
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate([20, 16, 12, 14, 14, 12, 12, 10, 52, 52], 1):
        tq.column_dimensions[CL(i)].width = w
    for i, s in enumerate(shops, 5):
        cell(tq, i, 1, s["shop"], font=BB); cell(tq, i, 2, s["direction"]); cell(tq, i, 3, s["current"], font=BLUE)
        cell(tq, i, 4, f"=SUMIF('Kế hoạch tuần'!$B$2:$B${last},A{i},'Kế hoạch tuần'!${TOTC}$2:${TOTC}${last})", font=GREEN)
        cell(tq, i, 5, f"=C{i}+D{i}", font=BB)
        cell(tq, i, 6, f"=SUMIF('Theo dõi thực tế'!$B$2:$B${last},A{i},'Theo dõi thực tế'!${DU}$2:${DU}${last})", font=GREEN)
        cell(tq, i, 7, f"=IF(D{i}=0,0,F{i}/D{i})", fmt="0%")
        cell(tq, i, 8, f"=COUNTIF('Kế hoạch tuần'!$B$2:$B${last},A{i})")
        cell(tq, i, 9, ", ".join(s["conflicts"]))
        cell(tq, i, 10, ", ".join(f'{m["ptype"]} {m["pct"]}%' for m in s["current_mix"]))
    e = 5 + len(shops)
    cell(tq, e, 1, "TỔNG", font=BB, fill=GREY)
    for c in [3, 4, 5, 6, 8]:
        cell(tq, e, c, f"=SUM({CL(c)}5:{CL(c)}{e - 1})", font=BB, fill=GREY)
    cell(tq, e, 7, f"=IF(D{e}=0,0,F{e}/D{e})", font=BB, fill=GREY, fmt="0%")
    tq.cell(row=e + 1, column=1, value="Listing hiện có: đếm từ CSV listing trong thư mục build/conflict (Everbee export).").font = ITAL
    tq.freeze_panes = "A5"

    # Check conflict
    cc = wb.create_sheet("Check conflict")
    header(cc, ["Shop build", "Dòng SP / loại SP", "Trạng thái", "Max %", "Chi tiết shop conflict", "Có trong plan?",
                "Nhu cầu thị trường (0–100)"], [20, 40, 13, 9, 64, 14, 14])
    r_i = 2
    for s in shops:
        items = [(r["line"], r["ptype"], row_conflict(r, rows, ctx), "Có", r.get("apparel")) for r in rows if r["shop"] == s["shop"]]
        items += [(b["line"], b["line"], dict(b, level="block"), "Không — chặn", s["direction"] == "Clothing") for b in s["blocked"]]
        for name, pt, c, inplan, app in items:
            lv = LEVEL_TXT[c["level"]]
            dm = payload["demand"].get("Apparel" if app else pt)
            for j, v in enumerate([s["shop"], name, lv, c["max"] / 100, hit_txt(c), inplan, dm if dm is not None else "—"], 1):
                cell(cc, r_i, j, v, fmt="0.0%" if j == 4 else None)
            cc.cell(row=r_i, column=3).fill = PatternFill("solid", fgColor=LEVEL_FILL[lv])
            r_i += 1
    cc.freeze_panes = "C2"; cc.auto_filter.ref = f"A1:G{max(r_i - 1, 1)}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


PLAN_CSS = """
.lp-kpis { display: grid; grid-template-columns: repeat(auto-fit, minmax(170px, 1fr)); gap: 1px; background: var(--line);
  border: 1px solid var(--line); border-radius: 6px; overflow: hidden; margin: 0 0 14px; }
.lp-kpi { background: var(--paper-raised); padding: 12px 14px; display: grid; gap: 2px; }
.lp-kpi .k { font-size: 11px; text-transform: uppercase; letter-spacing: .05em; color: var(--ink-soft); font-weight: 600; }
.lp-kpi .v { font-size: 22px; font-weight: 700; font-variant-numeric: tabular-nums; }
.lp-kpi .h { font-size: 12px; color: var(--ink-soft); }
.lp-cap { display: grid; grid-template-columns: repeat(var(--nw), minmax(68px, 1fr)); gap: 8px; overflow-x: auto; margin: 0 0 6px; }
.lp-wk .bars { position: relative; height: 84px; border-bottom: 1px solid var(--line); }
.lp-wk .band { position: absolute; left: 0; right: 0; background: var(--accent-soft); border-radius: 2px; }
.lp-wk .plan { position: absolute; left: 22%; width: 26%; bottom: 0; background: var(--accent); border-radius: 2px 2px 0 0; }
.lp-wk .done { position: absolute; left: 52%; width: 26%; bottom: 0; background: var(--warn); border-radius: 2px 2px 0 0; }
.lp-wk .t { display: flex; justify-content: space-between; font-size: 12px; margin-top: 3px; font-variant-numeric: tabular-nums; }
.lp-wk .d { font-size: 11px; color: var(--ink-soft); }
.lp-wk.now .t b { color: var(--accent); }
.lp-wk.off .t span { color: var(--critical); }
.lp-legend { display: flex; flex-wrap: wrap; gap: 14px; font-size: 12px; color: var(--ink-soft); margin: 0 0 14px; }
.lp-sw { display: inline-block; width: 10px; height: 10px; border-radius: 2px; margin-right: 5px; vertical-align: -1px; }
.lp-tabs { display: flex; flex-wrap: wrap; gap: 6px; margin: 0 0 12px; }
.lp-tab { border: 1px solid var(--line-strong); background: var(--paper-raised); color: var(--ink); border-radius: 999px;
  padding: 5px 12px; font: inherit; font-size: 13px; cursor: pointer; }
.lp-tab .n { color: var(--ink-soft); font-size: 12px; margin-left: 6px; font-variant-numeric: tabular-nums; }
.lp-tab[aria-selected="true"] { background: var(--accent); border-color: var(--accent); color: #fff; }
.lp-tab[aria-selected="true"] .n { color: #fff; }
.lp-head { display: grid; grid-template-columns: minmax(0, 1.2fr) minmax(0, 1fr); gap: 16px; margin: 0 0 12px; }
@media (max-width: 760px) { .lp-head { grid-template-columns: 1fr; } }
.lp-facts { display: flex; flex-wrap: wrap; gap: 4px 18px; font-size: 13px; color: var(--ink-soft); margin-top: 6px; }
.lp-facts b { color: var(--ink); font-variant-numeric: tabular-nums; }
.lp-chips { display: flex; flex-wrap: wrap; gap: 5px; margin: 4px 0 8px; }
.lp-chip { font-size: 12px; padding: 2px 8px; border-radius: 4px; background: var(--muted-soft); color: var(--ink-soft); }
.lp-chip.bad { background: var(--critical-soft); color: var(--critical); }
.lp-lbl { font-size: 11px; text-transform: uppercase; letter-spacing: .05em; color: var(--ink-soft); font-weight: 600; }
.lp-table td, .lp-table th { padding: 5px 6px; vertical-align: middle; }
.lp-table th.w, .lp-table td.w { text-align: center; }
.lp-table th.now, .lp-table td.now { background: var(--accent-soft); }
.lp-cell { display: grid; gap: 2px; justify-items: center; }
.lp-cell input { width: 44px; font: inherit; font-size: 13px; text-align: center; font-variant-numeric: tabular-nums;
  border: 1px solid var(--line-strong); border-radius: 3px; padding: 2px 0; background: var(--paper-raised); color: var(--ink); }
.lp-cell input.dn { font-size: 12px; border-style: dashed; color: var(--warn); }
.lp-cell input.dn.hit { border-style: solid; border-color: var(--good); color: var(--good); }
.lp-table .ln { font-weight: 600; }
.lp-table .pt { font-size: 12px; color: var(--ink-soft); }
.lp-table .tot { text-align: center; font-weight: 600; white-space: nowrap; font-variant-numeric: tabular-nums; }
.lp-table .tot small { display: block; font-weight: 400; color: var(--ink-soft); }
.lp-table td.note input { width: 100%; min-width: 180px; font: inherit; font-size: 12px; border: 1px solid transparent;
  border-radius: 3px; padding: 3px 6px; background: transparent; color: var(--ink); }
.lp-table td.note input:hover, .lp-table td.note input:focus { border-color: var(--line-strong); background: var(--paper-raised); }
.lp-table tfoot td { font-weight: 600; background: var(--muted-soft); }
.pill.lv-ok { background: var(--good-soft); color: var(--good); }
.pill.lv-light { background: var(--warn-soft); color: var(--warn); }
.pill.lv-block { background: var(--critical-soft); color: var(--critical); }
.pill.lv-wait { background: var(--accent-soft); color: var(--accent); }
.lp-del { background: none; border: none; color: var(--ink-soft); cursor: pointer; font-size: 16px; padding: 2px 6px; border-radius: 3px; }
.lp-del:hover { color: var(--critical); background: var(--critical-soft); }
.lp-add { display: grid; grid-template-columns: minmax(220px, 2fr) repeat(3, minmax(110px, 1fr)) auto; gap: 8px; align-items: end; margin: 14px 0 4px; }
@media (max-width: 760px) { .lp-add { grid-template-columns: 1fr 1fr; } }
.lp-add label { display: grid; gap: 4px; font-size: 12px; color: var(--ink-soft); }
.lp-add input, .lp-add select { font: inherit; font-size: 13.5px; padding: 7px 9px; border: 1px solid var(--line-strong);
  border-radius: 3px; background: var(--paper-raised); color: var(--ink); min-width: 0; }
.lp-status { font-size: 12.5px; color: var(--ink-soft); }
.lp-status.err { color: var(--critical); }
.lp-sched { margin: 0 0 16px; }
.lp-sched-head { display: flex; flex-wrap: wrap; gap: 8px 14px; align-items: center; margin: 0 0 8px; }
.lp-weeks { display: flex; flex-wrap: wrap; gap: 4px; }
.lp-weeks button { border: 1px solid var(--line-strong); background: var(--paper-raised); color: var(--ink); border-radius: 3px;
  padding: 3px 8px; font: inherit; font-size: 12px; cursor: pointer; }
.lp-weeks button[aria-pressed="true"] { background: var(--accent); border-color: var(--accent); color: #fff; }
.lp-table td.b5 { text-align: center; font-size: 12px; font-weight: 600; color: var(--accent); background: var(--accent-soft); min-width: 120px; }
.lp-table td.b5 small { display: block; font-weight: 400; color: var(--ink-soft); }
.lp-table td.b5.sh { box-shadow: inset 0 0 0 2px var(--accent); }
.lp-table td.b0 { text-align: center; color: var(--muted); }
.lp-table tr.today td { box-shadow: inset 0 -2px 0 var(--accent); }
.lp-bar { display: flex; flex-wrap: wrap; gap: 8px; align-items: center; justify-content: space-between; margin: 0 0 12px; }
.lp-bar .btns { display: flex; gap: 8px; flex-wrap: wrap; }
"""

PLAN_PANE = """
      <div class="pane" id="bc-pane-listingplan" hidden>
        <p class="reverse-intro">
          Target số listing <strong>mới</strong> mỗi shop build phải up tới hết <strong>__PLAN_END__</strong>, tách theo tuần và
          dòng sản phẩm. VowgiftNina đi <strong>handmade</strong>, YourStoryGiftsArt đi <strong>clothing</strong>, các shop còn lại
          đi <strong>custom gift POD</strong>. Niche không giới hạn trong một shop. Mỗi ngày up <strong>__PER_DAY_MIN__–__PER_DAY_MAX__ listing</strong>
          (làm T2–T6, riêng tuần đầu tháng làm thêm T7) chia thành <strong>lượt 5 listing, mỗi lượt thuộc 1 dòng sản phẩm của 1 shop</strong>; mỗi shop tối đa 1 lượt/ngày.
          Cùng ngày nhiều shop có thể up cùng một dòng SP (dùng chung mẫu, khác shop). Khoảng 90% số ngày up 20 listing (4 lượt),
          ~10% số ngày up 15 listing (3 lượt) dồn vào đầu W1 để mở dòng mới. Mỗi shop tập trung <strong>4–6 dòng SP</strong>,
          riêng YourStoryGiftsArt <strong>3 dòng</strong>. Target mỗi dòng mỗi tuần là bội số của 5 (= số lượt × 5); lịch theo ngày bên dưới tự chia lại khi sửa target. Dòng Giáng Sinh dồn
          vào các tuần đầu vì listing mới cần 2–4 tuần mới có view.
          <strong>Conflict</strong> = % listing của shop conflict cùng loại SP tại tuần bắt đầu của dòng (shop conflict cũng là
          shop build thì tính cả listing nó sẽ up theo plan): <span class="pill lv-block">Chặn</span> ≥ 10%,
          <span class="pill lv-wait">Chờ mở</span> hiện ≥ 10% nhưng xuống dưới 10% khi tới tuần bắt đầu,
          <span class="pill lv-light">Nhẹ</span> dưới 10% (up được, tránh trùng mẫu của shop đó). Quần áo so theo dòng con
          (Sweatshirt, Hoodie, T-Shirt…). Sửa ô nào là tự lưu vào file trên máy chạy tool.
        </p>
        <div class="lp-kpis" id="lp-kpis"></div>
        <div class="lp-cap" id="lp-cap"></div>
        <div class="lp-legend">
          <span><span class="lp-sw" style="background:var(--accent-soft)"></span>Dải __PER_DAY_MIN__–__PER_DAY_MAX__ mẫu/ngày</span>
          <span><span class="lp-sw" style="background:var(--accent)"></span>Target</span>
          <span><span class="lp-sw" style="background:var(--warn)"></span>Đã up</span>
        </div>
        <div class="lp-sched">
          <div class="lp-sched-head"><span class="lp-lbl">Lịch up theo ngày</span><span class="lp-weeks" id="lp-sched-weeks"></span></div>
          <div class="table-scroll"><table class="lp-table" id="lp-sched"></table></div>
          <p class="chart-caveat" id="lp-sched-note"></p>
        </div>
        <div class="lp-bar">
          <span class="lp-status" id="lp-status"></span>
          <span class="btns">
            <a class="btn ghost" id="lp-xlsx" href="/build-conflict/plan.xlsx">Tải Excel</a>
            <button type="button" class="btn ghost" id="lp-reset">Khôi phục plan gốc</button>
          </span>
        </div>
        <div class="lp-tabs" role="tablist" id="lp-tabs"></div>
        <div class="lp-head" id="lp-head"></div>
        <div class="table-scroll"><table class="lp-table" id="lp-grid" style="min-width:1180px"></table></div>
        <form class="lp-add" id="lp-add" autocomplete="off">
          <label>Thêm dòng sản phẩm cho shop này
            <input id="lp-add-name" type="text" placeholder="vd: Custom Pet Blanket">
          </label>
          <label>Mùa vụ
            <select id="lp-add-season"><option value="E">Evergreen</option><option value="X">Giáng Sinh</option></select>
          </label>
          <label>Tuần bắt đầu <select id="lp-add-start"></select></label>
          <label>Lượt 5 listing / tuần <input id="lp-add-per" type="number" min="0" max="6" value="1"></label>
          <button class="btn" id="lp-add-btn" type="submit" disabled>Thêm</button>
        </form>
        <p class="chart-caveat" id="lp-detect">&nbsp;</p>
      </div>
"""

PLAN_JS = r"""
<script>
(function () {
  var bc = JSON.parse(document.getElementById('build-conflict-data').textContent);
  if (!bc.available || !bc.listing_plan) return;
  var P = bc.listing_plan, W = P.weeks, NW = W.length;
  var SHOPS = P.shops.map(function (s) { return s.shop; });
  var SHOP = {}; P.shops.forEach(function (s) { SHOP[s.shop] = s; });
  var CONF = {}; bc.planner.build_list.forEach(function (b) { CONF[b.build_shop] = b.conflicts; });
  var PT = bc.planner.ptype_rules.map(function (r) { return [r[0], new RegExp(r[1])]; });
  var AP = P.apparel_rules.map(function (r) { return [r[0], new RegExp(r[1])]; });
  var LV = { ok: 'OK', light: 'Nhẹ', block: 'Chặn', wait: 'Chờ mở' };
  var rows = P.rows.map(function (r) { return Object.assign({}, r); });
  var LOCAL_KEY = 'buildconf-listing-plan-v1';
  var serverOk = location.protocol !== 'file:';
  if (!serverOk) {
    try { var loc = JSON.parse(localStorage.getItem(LOCAL_KEY) || 'null'); if (loc && Array.isArray(loc.rows)) rows = loc.rows; } catch (e) {}
  }
  var current = SHOPS[0];
  try { var tab = localStorage.getItem('buildconf-listing-plan-tab'); if (tab && SHOP[tab]) current = tab; } catch (e) {}
  var statusEl = document.getElementById('lp-status');

  function esc(s) { return String(s == null ? '' : s).replace(/[&<>"']/g, function (c) { return { '&': '&amp;', '<': '&lt;', '>': '&gt;', '"': '&quot;', "'": '&#39;' }[c]; }); }
  function sum(a) { return (a || []).reduce(function (x, y) { return x + (Number(y) || 0); }, 0); }
  function sumTo(a, n) { return sum((a || []).slice(0, n)); }
  function setStatus(t, err) { statusEl.textContent = t; statusEl.className = 'lp-status' + (err ? ' err' : ''); }
  function shopRows(s) { return rows.filter(function (r) { return r.shop === s; }).sort(function (a, b) { return a.order - b.order; }); }
  function nowWeek() {
    var t = new Date(); t.setHours(0, 0, 0, 0);
    for (var i = 0; i < NW; i++) {
      var a = new Date(W[i].start + 'T00:00:00'), b = new Date(W[i].end + 'T00:00:00'); b.setDate(b.getDate() + 1);
      if (t >= a && t < b) return i;
    }
    return t < new Date(W[0].start + 'T00:00:00') ? -1 : NW;
  }
  var NOW = nowWeek();

  // Cùng công thức với listing_plan.row_conflict (Python) — % tại tuần bắt đầu của dòng.
  function conflictOf(shop, key, apparel, start) {
    var src = apparel ? P.ctx.app_n : P.ctx.ptype_n, hits = [], maxNow = 0;
    (CONF[shop] || []).forEach(function (c) {
      var ck = c.toLowerCase(), baseN = (src[ck] || {})[key] || 0, baseT = P.ctx.total[ck] || 0;
      var cr = rows.filter(function (r) { return r.shop.toLowerCase() === ck; });
      var addT = 0, addN = 0, planned = false;
      cr.forEach(function (r) {
        addT += sumTo(r.t, start - 1);
        if (r.ptype === key && !!r.apparel === !!apparel) { addN += sumTo(r.t, start - 1); if (sum(r.t) > 0) planned = true; }
      });
      var now = baseT ? Math.round(baseN / baseT * 1000) / 10 : 0;
      var at = baseT + addT ? Math.round((baseN + addN) / (baseT + addT) * 1000) / 10 : 0;
      maxNow = Math.max(maxNow, now);
      if (at > 0 || now > 0 || planned) hits.push({ shop: c, pct: at, now: now, planned: planned });
    });
    hits.sort(function (a, b) { return b.pct - a.pct; });
    var mx = hits.length ? hits[0].pct : 0;
    var level = mx >= P.block_pct ? 'block' : maxNow >= P.block_pct ? 'wait' : (mx > 0 || hits.some(function (h) { return h.planned; })) ? 'light' : 'ok';
    return { level: level, max: mx, hits: hits };
  }
  function pill(c) {
    var tip = c.hits.map(function (h) { return h.shop + ' ' + h.pct + '%' + (h.now !== h.pct ? ' (hiện ' + h.now + '%)' : '') + (h.planned ? ', đã lên plan' : ''); }).join('; ') || 'Không shop conflict nào bán loại này';
    return '<span class="pill lv-' + c.level + '" title="' + esc(tip) + '">' + LV[c.level] + (c.max ? ' ' + c.max + '%' : '') + '</span>';
  }

  // ── Lưu ──
  var saveTimer = null, saving = false, dirty = false;
  function scheduleSave() { dirty = true; clearTimeout(saveTimer); saveTimer = setTimeout(save, 700); setStatus('Chưa lưu…'); }
  function save() {
    if (!dirty) return Promise.resolve();
    if (!serverOk) {
      try { localStorage.setItem(LOCAL_KEY, JSON.stringify({ rows: rows })); dirty = false; setStatus('Đã lưu trên trình duyệt này (mở report qua server để lưu ra file)'); } catch (e) { setStatus('Không lưu được', true); }
      return Promise.resolve();
    }
    if (saving) { clearTimeout(saveTimer); saveTimer = setTimeout(save, 400); return Promise.resolve(); }
    saving = true; dirty = false; setStatus('Đang lưu…');
    return fetch('/build-conflict/plan', { method: 'POST', headers: { 'Content-Type': 'application/json' }, body: JSON.stringify({ rows: rows }) })
      .then(function (r) { if (!r.ok) throw new Error('HTTP ' + r.status); return r.json(); })
      .then(function (j) { setStatus('Đã lưu ' + j.saved_at.replace('T', ' ') + ' vào ' + P.plan_file); },
            function (e) { dirty = true; setStatus('Chưa lưu được (' + e.message + '). Kiểm tra server còn chạy rồi sửa lại một ô.', true); })
      .then(function () { saving = false; });
  }
  setStatus(P.saved_at ? 'Plan đã lưu lúc ' + P.saved_at.replace('T', ' ') : 'Đang xem plan gốc — sửa ô nào sẽ lưu ra ' + P.plan_file);

  // ── Tổng quan + công suất ──
  function renderTop() {
    var wt = [], wd = [];
    for (var i = 0; i < NW; i++) { wt[i] = 0; wd[i] = 0; }
    rows.forEach(function (r) { for (var i = 0; i < NW; i++) { wt[i] += Number(r.t[i]) || 0; wd[i] += Number(r.d[i]) || 0; } });
    var t = sum(wt), d = sum(wd), cur = P.shops.reduce(function (x, s) { return x + s.current; }, 0), due = sumTo(wt, Math.max(0, Math.min(NOW, NW)));
    var lo = sum(W.map(function (w) { return w.days * P.per_day_min; })), hi = sum(W.map(function (w) { return w.days * P.per_day_max; }));
    document.getElementById('lp-kpis').innerHTML = [
      ['Target listing mới', t, 'dải công suất ' + lo + '–' + hi],
      ['Đã up', d, t ? Math.round(d / t * 100) + '% target' : '—'],
      ['Tổng listing ' + W[NW - 1].to, cur + t, cur + ' hiện có + ' + t + ' mới'],
      ['Nên đạt tới tuần này', due, NOW <= 0 ? 'chưa qua tuần nào' : (d >= due ? 'đang vượt ' : 'đang thiếu ') + Math.abs(d - due)],
      ['Tuần này', NOW >= 0 && NOW < NW ? wt[NOW] : '—', NOW >= 0 && NOW < NW ? W[NOW].code + ' · ' + W[NOW].from + '–' + W[NOW].to : (NOW < 0 ? 'Bắt đầu ' + W[0].code + ' ngày ' + W[0].from : 'Đã hết kỳ plan')]
    ].map(function (k) { return '<div class="lp-kpi"><span class="k">' + esc(k[0]) + '</span><span class="v">' + esc(k[1]) + '</span><span class="h">' + esc(k[2]) + '</span></div>'; }).join('');
    var mx = Math.max.apply(null, W.map(function (w, i) { return Math.max(w.days * P.per_day_max, wt[i], wd[i]); }).concat([1])), h = 84;
    var cap = document.getElementById('lp-cap'); cap.style.setProperty('--nw', NW);
    cap.innerHTML = W.map(function (w, i) {
      var a = w.days * P.per_day_min, b = w.days * P.per_day_max, off = wt[i] < a || wt[i] > b;
      return '<div class="lp-wk' + (i === NOW ? ' now' : '') + (off ? ' off' : '') + '" title="' + w.code + ' ' + w.from + '–' + w.to + ': target ' + wt[i] + ' (dải ' + a + '–' + b + '), đã up ' + wd[i] + '">' +
        '<div class="bars"><div class="band" style="bottom:' + (a / mx * h) + 'px;height:' + ((b - a) / mx * h) + 'px"></div>' +
        '<div class="plan" style="height:' + (wt[i] / mx * h) + 'px"></div><div class="done" style="height:' + (wd[i] / mx * h) + 'px"></div></div>' +
        '<div class="t"><b>' + w.code + '</b><span>' + wd[i] + '/' + wt[i] + '</span></div><div class="d">' + w.from + '–' + w.to + ' · ' + w.days + ' ngày</div></div>';
    }).join('');
  }
  function renderTabs() {
    document.getElementById('lp-tabs').innerHTML = SHOPS.map(function (s) {
      var t = 0; shopRows(s).forEach(function (r) { t += sum(r.t); });
      return '<button type="button" role="tab" class="lp-tab" aria-selected="' + (s === current) + '" data-shop="' + esc(s) + '">' + esc(s) + '<span class="n">' + t + '</span></button>';
    }).join('');
  }
  document.getElementById('lp-tabs').addEventListener('click', function (e) {
    var b = e.target.closest('.lp-tab'); if (!b) return;
    current = b.getAttribute('data-shop');
    try { localStorage.setItem('buildconf-listing-plan-tab', current); } catch (err) {}
    renderTabs(); renderShop(); refreshDetect();
  });
  function renderShop() {
    var s = SHOP[current], rs = shopRows(current), t = 0, d = 0;
    rs.forEach(function (r) { t += sum(r.t); d += sum(r.d); });
    document.getElementById('lp-head').innerHTML =
      '<div><span class="pill role-build">' + esc(s.direction) + '</span>' +
      '<div class="lp-facts"><span>Hiện có <b>' + s.current + '</b> listing</span><span' + (rs.length < s.lines[0] || rs.length > s.lines[1] ? ' style="color:var(--critical)"' : '') + '>Dòng SP <b>' + rs.length + '</b> (mục tiêu ' + (s.lines[0] === s.lines[1] ? s.lines[0] : s.lines[0] + '–' + s.lines[1]) + ')</span><span>Target mới <b>' + t + '</b></span><span>Đã up <b>' + d + '</b></span><span>Dự kiến ' + W[NW - 1].to + ' <b>' + (s.current + t) + '</b></span></div>' +
      '<div class="lp-facts">Đang bán nhiều nhất: ' + esc(s.current_mix.map(function (m) { return m.ptype + ' ' + m.pct + '%'; }).join(' · ') || '—') + '</div></div>' +
      '<div><div class="lp-lbl">Shop conflict</div><div class="lp-chips">' + s.conflicts.map(function (c) { return '<span class="lp-chip">' + esc(c) + '</span>'; }).join('') + '</div>' +
      '<div class="lp-lbl">Không up (conflict ≥ ' + P.block_pct + '%)</div><div class="lp-chips">' + (s.blocked.map(function (b) {
        return '<span class="lp-chip bad" title="' + esc(b.hits.map(function (h) { return h.shop + ' ' + h.pct + '%'; }).join(', ')) + '">' + esc(b.line) + ' ' + b.max + '%</span>'; }).join('') || '<span class="lp-chip">—</span>') + '</div></div>';
    var head = '<thead><tr><th>Dòng sản phẩm</th><th>Conflict</th>' + W.map(function (w, i) { return '<th class="w' + (i === NOW ? ' now' : '') + '">' + w.code + '<br><span style="font-weight:400">' + w.from + '</span></th>'; }).join('') + '<th class="w">Tổng</th><th>Ghi chú</th><th></th></tr></thead>';
    if (!rs.length) { document.getElementById('lp-grid').innerHTML = head + '<tbody><tr><td colspan="' + (NW + 5) + '" class="empty-state">Shop này chưa có dòng sản phẩm nào. Thêm dòng đầu tiên bằng ô bên dưới.</td></tr></tbody>'; return; }
    var body = rs.map(function (r) {
      return '<tr data-id="' + esc(r.id) + '"><td><div class="ln">' + esc(r.line) + '</div><div class="pt">' + esc(r.ptype) + ' · ' + (r.season === 'X' ? 'Giáng Sinh' : 'Evergreen') + (r.start > 1 ? ' · từ W' + r.start : '') + '</div></td>' +
        '<td>' + pill(conflictOf(r.shop, r.ptype, r.apparel, r.start)) + '</td>' +
        W.map(function (w, i) {
          var tv = Number(r.t[i]) || 0, dv = Number(r.d[i]) || 0;
          return '<td class="w' + (i === NOW ? ' now' : '') + '"><div class="lp-cell">' +
            '<input type="number" min="0" max="99" step="' + P.batch + '" data-k="t" data-i="' + i + '" value="' + tv + '" aria-label="Target ' + w.code + '"' + (tv % P.batch ? ' style="border-color:var(--critical)" title="Không chia hết cho ' + P.batch + '"' : '') + '>' +
            '<input type="number" min="0" max="99" class="dn' + (tv && dv >= tv ? ' hit' : '') + '" data-k="d" data-i="' + i + '" value="' + (dv || '') + '" placeholder="·" aria-label="Đã up ' + w.code + '"></div></td>';
        }).join('') +
        '<td class="tot">' + sum(r.t) + '<small>đã up ' + sum(r.d) + '</small></td>' +
        '<td class="note"><input type="text" data-k="note" value="' + esc(r.note) + '" aria-label="Ghi chú"></td>' +
        '<td><button type="button" class="lp-del" title="Xoá dòng">×</button></td></tr>';
    }).join('');
    var ft = W.map(function (w, i) { var a = 0, b = 0; rs.forEach(function (r) { a += Number(r.t[i]) || 0; b += Number(r.d[i]) || 0; }); return '<td class="w">' + a + '<br><span style="color:var(--warn);font-weight:400">' + b + '</span></td>'; }).join('');
    document.getElementById('lp-grid').innerHTML = head + '<tbody>' + body + '</tbody><tfoot><tr><td>Tổng shop</td><td></td>' + ft + '<td class="tot">' + t + '<small>đã up ' + d + '</small></td><td colspan="2"></td></tr></tfoot>';
  }
  var grid = document.getElementById('lp-grid');
  grid.addEventListener('input', function (e) {
    var el = e.target, tr = el.closest('tr[data-id]'); if (!tr || !el.dataset.k) return;
    var r = rows.filter(function (x) { return x.id === tr.getAttribute('data-id'); })[0]; if (!r) return;
    if (el.dataset.k === 'note') { r.note = el.value; scheduleSave(); return; }
    var i = Number(el.dataset.i), v = Math.max(0, Math.min(99, parseInt(el.value, 10) || 0));
    r[el.dataset.k] = r[el.dataset.k].slice(); r[el.dataset.k][i] = v;
    if (el.dataset.k === 'd') el.classList.toggle('hit', !!r.t[i] && v >= r.t[i]);
    var c = tr.querySelector('.tot'); if (c) c.innerHTML = sum(r.t) + '<small>đã up ' + sum(r.d) + '</small>';
    renderTop(); renderTabs(); if (el.dataset.k === 't') renderSchedule(); scheduleSave();
  });
  grid.addEventListener('change', function (e) { if (e.target.dataset.k === 't' || e.target.dataset.k === 'd') { renderShop(); refreshDetect(); } });
  grid.addEventListener('click', function (e) {
    var b = e.target.closest('.lp-del'); if (!b) return;
    var id = b.closest('tr[data-id]').getAttribute('data-id');
    var r = rows.filter(function (x) { return x.id === id; })[0];
    if (!r || !confirm('Xoá dòng "' + r.line + '" khỏi ' + r.shop + '?')) return;
    rows = rows.filter(function (x) { return x.id !== id; });
    renderAll(); scheduleSave();
  });
  document.getElementById('lp-reset').addEventListener('click', function () {
    if (!confirm('Bỏ mọi chỉnh sửa và quay về plan gốc? File plan hiện tại sẽ được giữ lại dạng .bak.')) return;
    if (!serverOk) { try { localStorage.removeItem(LOCAL_KEY); } catch (e) {} location.reload(); return; }
    fetch('/build-conflict/plan/reset', { method: 'POST' }).then(function () { location.reload(); }, function () { setStatus('Không khôi phục được — server có đang chạy không?', true); });
  });
  document.getElementById('lp-xlsx').addEventListener('click', function (e) {
    if (!serverOk) { e.preventDefault(); setStatus('Tải Excel cần mở report qua server (python analyze_shop_gift_fit.py).', true); return; }
    if (dirty || saving) { e.preventDefault(); var href = this.href; clearTimeout(saveTimer); save().then(function () { location.href = href; }); }
  });

  // ── Thêm dòng ──
  var nameEl = document.getElementById('lp-add-name'), detEl = document.getElementById('lp-detect'), addBtn = document.getElementById('lp-add-btn');
  document.getElementById('lp-add-start').innerHTML = W.map(function (w, i) { return '<option value="' + (i + 1) + '"' + (i === Math.max(0, Math.min(NOW, NW - 1)) ? ' selected' : '') + '>' + w.code + ' (' + w.from + ')</option>'; }).join('');
  function classify(name, rules) { var t = name.toLowerCase(); for (var i = 0; i < rules.length; i++) if (rules[i][1].test(t)) return rules[i][0]; return null; }
  function detect(name, start) {
    var apparel = SHOP[current].direction === 'Clothing' || classify(name, PT) === 'Apparel';
    var key = apparel ? (classify(name, AP) || 'Apparel khác') : (classify(name, PT) || 'Other');
    return { key: key, apparel: apparel, conflict: conflictOf(current, key, apparel, start) };
  }
  function refreshDetect() {
    var n = nameEl.value.trim();
    if (!n) { detEl.innerHTML = '&nbsp;'; addBtn.disabled = true; return; }
    var r = detect(n, Number(document.getElementById('lp-add-start').value)), c = r.conflict;
    var dup = shopRows(current).some(function (x) { return x.line.toLowerCase() === n.toLowerCase(); });
    var msg = 'Loại SP: <strong>' + esc(r.key) + '</strong> ' + pill(c) + ' ';
    if (r.key === 'Other') msg += 'Không khớp loại SP nào nên chưa check được conflict — kiểm tra tay trước khi up.';
    else if (c.level === 'block') msg += 'Không thêm được: ' + esc(c.hits[0].shop) + ' có ' + c.hits[0].pct + '% listing loại này.';
    else if (c.level === 'light') msg += 'Up được, tránh trùng mẫu của ' + esc(c.hits.map(function (h) { return h.shop + ' ' + h.pct + '%'; }).join(', ')) + '.';
    else if (c.level === 'wait') msg += 'Hiện đang vướng conflict nhưng xuống dưới ' + P.block_pct + '% tại tuần bắt đầu đã chọn.';
    else msg += 'Không shop conflict nào bán loại này.';
    if (dup) msg += ' Dòng này đã có trong shop.';
    else if (shopRows(current).length >= SHOP[current].lines[1]) msg += ' Shop đã đủ ' + SHOP[current].lines[1] + ' dòng tập trung — nên bỏ bớt 1 dòng trước khi thêm.';
    detEl.innerHTML = msg;
    addBtn.disabled = c.level === 'block' || dup;
  }
  nameEl.addEventListener('input', refreshDetect);
  document.getElementById('lp-add-start').addEventListener('change', refreshDetect);
  document.getElementById('lp-add').addEventListener('submit', function (e) {
    e.preventDefault(); if (addBtn.disabled) return;
    var n = nameEl.value.trim(), start = Number(document.getElementById('lp-add-start').value), r = detect(n, start);
    var per = Math.max(0, Math.min(6, parseInt(document.getElementById('lp-add-per').value, 10) || 0));
    var order = shopRows(current).reduce(function (m, x) { return Math.max(m, x.order || 0); }, 0) + 1;
    rows.push({ id: current + '-' + Date.now().toString(36), shop: current, order: order, line: n, ptype: r.key, apparel: r.apparel,
      season: document.getElementById('lp-add-season').value, start: start, note: 'Thêm trên trang',
      t: W.map(function (w, i) { return i + 1 >= start ? Math.min(per, w.days) * P.batch : 0; }), d: W.map(function () { return 0; }) });
    nameEl.value = ''; renderAll(); refreshDetect(); scheduleSave();
  });

  // ── Lịch theo ngày: cùng thuật toán listing_plan.daily_schedule ──
  var schedWeek = Math.max(0, Math.min(NOW, NW - 1));
  function lineBatches(rs, wi) {  // = listing_plan._line_batches
    var q = rs.map(function (r) { var t = Number(r.t[wi]) || 0, b = []; for (var k = 0; k < t; k += P.batch) b.push({ line: r.line, ptype: r.ptype, n: Math.min(P.batch, t - k) }); return b; });
    var out = [], more = true;
    while (more) { more = false; q.forEach(function (b) { if (b.length) { out.push(b.shift()); more = true; } }); }
    return out;
  }
  function dailySchedule(wi) {
    var order = SHOPS, queue = {}, need = {}, days = W[wi].dates, out = [];
    order.forEach(function (s) { queue[s] = lineBatches(shopRows(s), wi); need[s] = queue[s].length; });
    days.forEach(function (d, i) {
      var left = days.length - i, totalNeed = 0; order.forEach(function (s) { totalNeed += need[s]; });
      var later = days.slice(i + 1).reduce(function (x, y) { return x + (y.cap || 0); }, 0);
      var cap = Math.max(d.cap || 0, totalNeed - later, d.cap ? 0 : Math.ceil(totalNeed / left));
      var must = order.filter(function (s) { return need[s] >= left; });
      var rest = order.filter(function (s) { return need[s] > 0 && must.indexOf(s) < 0; }).sort(function (a, b) { return need[b] - need[a] || order.indexOf(a) - order.indexOf(b); });
      var pick = must.concat(rest.slice(0, Math.max(0, cap - must.length)));
      var items = {}, today = {};
      order.filter(function (s) { return pick.indexOf(s) >= 0; }).forEach(function (s) {
        var q = queue[s], j = 0;
        for (var k = 0; k < q.length; k++) if (today[q[k].ptype]) { j = k; break; }
        var b = q.splice(j, 1)[0]; need[s] -= 1; today[b.ptype] = (today[b.ptype] || 0) + 1; items[s] = b;
      });
      out.push({ label: d.label, iso: d.iso, items: items, shared: today });
    });
    var left = 0; order.forEach(function (s) { queue[s].forEach(function (b) { left += b.n; }); });
    return { days: out, left: left };
  }
  function renderSchedule() {
    document.getElementById('lp-sched-weeks').innerHTML = W.map(function (w, i) {
      return '<button type="button" data-w="' + i + '" aria-pressed="' + (i === schedWeek) + '">' + w.code + ' · ' + w.from + '</button>';
    }).join('');
    var sc = dailySchedule(schedWeek), today = new Date().toISOString().slice(0, 10);
    var head = '<thead><tr><th>Ngày</th>' + SHOPS.map(function (s) { return '<th class="w">' + esc(s) + '</th>'; }).join('') + '<th class="w">Số lượt</th><th class="w">Listing</th></tr></thead>';
    var tot = {}; SHOPS.forEach(function (s) { tot[s] = 0; });
    var body = sc.days.map(function (d) {
      var n = 0, k = 0;
      var cells = SHOPS.map(function (s) {
        var b = d.items[s]; if (!b) return '<td class="b0">·</td>';
        tot[s] += b.n; n += b.n; k++;
        var sh = d.shared[b.ptype] > 1;
        return '<td class="b5' + (sh ? ' sh' : '') + '" title="' + esc(b.ptype + (sh ? ' — shop khác cũng up loại này hôm nay, dùng chung mẫu' : '')) + '">' + esc(b.line) + '<small>' + b.n + ' listing</small></td>';
      }).join('');
      return '<tr' + (d.iso === today ? ' class="today"' : '') + '><td>' + esc(d.label) + '</td>' + cells + '<td class="w">' + k + '</td><td class="w"><strong>' + n + '</strong></td></tr>';
    }).join('');
    var all = 0; SHOPS.forEach(function (s) { all += tot[s]; });
    document.getElementById('lp-sched').innerHTML = head + '<tbody>' + body + '</tbody><tfoot><tr><td>Tổng tuần</td>' + SHOPS.map(function (s) { return '<td class="w">' + tot[s] + '</td>'; }).join('') + '<td></td><td class="w">' + all + '</td></tr></tfoot>';
    var notes = [];
    if (sc.left) notes.push(sc.left + ' listing chưa xếp được vì shop có nhiều lượt hơn số ngày trong tuần (mỗi shop tối đa 1 lượt 5 listing/ngày) — giảm target tuần của shop đó.');
    var odd = rows.filter(function (r) { return (Number(r.t[schedWeek]) || 0) % P.batch; });
    if (odd.length) notes.push('Target không chia hết cho ' + P.batch + ': ' + odd.map(function (r) { return r.shop + ' › ' + r.line; }).join(', ') + ' — lượt cuối của dòng đó sẽ ít hơn ' + P.batch + ' listing.');
    var dayN = function (d) { var n = 0; for (var k in d.items) n += d.items[k].n; return n; };
    var over = sc.days.filter(function (d) { return dayN(d) > P.per_day_max; }), under = sc.days.filter(function (d) { var n = dayN(d); return n && n < P.per_day_min; });
    if (over.length) notes.push(over.length + ' ngày vượt ' + P.per_day_max + ' listing.');
    if (under.length) notes.push(under.length + ' ngày dưới ' + P.per_day_min + ' listing.');
    document.getElementById('lp-sched-note').textContent = notes.join(' ') || 'Mỗi ô = 1 lượt 5 listing của dòng SP đó cho shop đó. Ô viền đậm = nhiều shop cùng up loại SP này trong ngày, dùng chung mẫu.';
  }
  document.getElementById('lp-sched-weeks').addEventListener('click', function (e) {
    var b = e.target.closest('button[data-w]'); if (!b) return; schedWeek = Number(b.getAttribute('data-w')); renderSchedule();
  });

  function renderAll() { renderTop(); renderTabs(); renderShop(); renderSchedule(); }
  renderAll();
  window.addEventListener('beforeunload', function (e) { if (dirty || saving) { save(); e.preventDefault(); e.returnValue = ''; } });
})();
</script>
"""
